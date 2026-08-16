"""
Office documents to PDF, via LibreOffice.

The conversion shells out to LibreOffice, so most of this suite skips when it
is absent — but the two things that must hold whether or not it is installed
are tested unconditionally: an unsupported extension is refused by name, and a
missing LibreOffice produces an explanatory error rather than a crash or an
empty file that looks like success.

Where the converter does run, the output is opened and read back. A PDF that
is the right size but does not contain the source text is not a conversion.
"""
import io

import pytest
from docx import Document as Docx
from openpyxl import Workbook

from docintel.pdf import convert
from docintel.pdf.engine import PDFEngineError
from docintel.pdf.text import extract

MARKER = "Quarterly revenue reconciliation"

office_available = convert._office_binary() is not None
needs_office = pytest.mark.skipif(
    not office_available, reason="LibreOffice is not installed",
)


def docx_bytes() -> bytes:
    document = Docx()
    document.add_heading("Board Paper", level=1)
    document.add_paragraph(MARKER)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def xlsx_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet["A1"] = MARKER
    sheet["A2"] = "North"
    sheet["B2"] = 128400
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


# --- true with or without LibreOffice ------------------------------------

def test_unsupported_extension_is_refused_by_name():
    with pytest.raises(PDFEngineError) as caught:
        convert.to_pdf(b"data", "exe")

    message = str(caught.value)
    assert ".exe" in message
    assert "docx" in message      # the error lists what would work


def test_missing_libreoffice_explains_itself(monkeypatch):
    """The failure must name the missing dependency, not just fail."""
    monkeypatch.setattr(convert, "_office_binary", lambda: None)

    with pytest.raises(PDFEngineError) as caught:
        convert.office_to_pdf(docx_bytes(), "docx")

    assert "LibreOffice" in str(caught.value)


def test_text_and_csv_do_not_need_libreoffice(monkeypatch):
    """These paths are pure Python and must keep working regardless."""
    monkeypatch.setattr(convert, "_office_binary", lambda: None)

    pdf = convert.to_pdf(b"hello from plain text", "txt", title="Note")
    assert pdf.startswith(b"%PDF")
    assert "hello from plain text" in " ".join(p.text for p in extract(pdf))


# --- only with LibreOffice -----------------------------------------------

@needs_office
def test_docx_becomes_a_pdf_carrying_its_text():
    pdf = convert.to_pdf(docx_bytes(), "docx")

    assert pdf.startswith(b"%PDF")
    pages = extract(pdf)
    assert pages, "the PDF has no pages"
    assert MARKER.lower() in " ".join(p.text for p in pages).lower()


@needs_office
def test_xlsx_becomes_a_pdf_carrying_its_cells():
    pdf = convert.to_pdf(xlsx_bytes(), "xlsx")

    assert pdf.startswith(b"%PDF")
    text = " ".join(p.text for p in extract(pdf))
    assert MARKER.lower() in text.lower()
    assert "128400" in text.replace(",", "").replace(" ", "")


@needs_office
def test_mislabelled_text_is_converted_rather_than_rejected():
    """LibreOffice trusts content over the extension, and that is fine.

    A text file named .docx comes back as a PDF of that text. Documenting it
    because the obvious expectation is a failure: what actually matters is
    that the output is a real PDF and not a silently empty one.
    """
    pdf = convert.to_pdf(b"this is not a docx at all", "docx")

    assert pdf.startswith(b"%PDF")
    assert "this is not a docx at all" in " ".join(p.text for p in extract(pdf))


@needs_office
def test_output_is_never_silently_empty():
    """Whatever the input, a returned PDF must be openable and have pages.

    The converter reports failure by raising; it must never return a
    zero-byte or pageless file that downstream code treats as a success.
    """
    for raw in (b"plain text", b"\x00\x01\x02binary noise" * 100):
        try:
            pdf = convert.to_pdf(raw, "docx")
        except PDFEngineError:
            continue          # refusing is an acceptable outcome
        assert pdf.startswith(b"%PDF")
        assert len(extract(pdf)) >= 1
